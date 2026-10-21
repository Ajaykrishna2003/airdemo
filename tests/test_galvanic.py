import importlib.util
import json
from pathlib import Path
import threading
import unittest
from urllib.request import Request, urlopen
from urllib.error import HTTPError

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('galvanic_server',ROOT/'backend/server.py')
server=importlib.util.module_from_spec(spec)
spec.loader.exec_module(server)

class GalvanicTests(unittest.TestCase):
    def setUp(self):
        self.db=object.__new__(server.Database)
        self.row=['TEST']+[None]*85
        self.row[80:86]=['Titanium alloy','None','Titanium alloy','None',None,None]
        self.db.by_part={'test':self.row}
        self.db.bonding_pair_galvanic=[['Titanium alloy','None','Titanium alloy','None','No','Bonding pair note']]
        self.db.pair_galvanic=[['Titanium alloy','None','Titanium alloy','None','Yes','Pair note']]
        self.db.galvanic=[['Titanium alloy','None','Aluminum alloy','No','Bolt note']]
        self.db.nut_galvanic=[['Titanium alloy','None','Aluminum alloy','TSA','No','Nut TSA note'],
                              ['Titanium alloy','None','Aluminum alloy','CCC','Yes','Nut CCC note']]
        self.db.bonding_galvanic=[['Titanium alloy','None','Aluminum alloy','Yes','Bonding bolt note']]
        self.db.bonding_nut_galvanic=[['Titanium alloy','None','Aluminum alloy','CCC','No']]
        self.item=dict(material='Aluminum alloy',finish='CCC',nutSide=True)
        self.query=dict(part='TEST',mode='no',substrates=[self.item])

    def test_bolt_ignores_finish_and_nut_uses_finish(self):
        status,result=self.db.compatibility(self.query)
        self.assertEqual(status,200)
        bolt,nut,extra=result['results']
        self.assertEqual(bolt['result'],'No')
        self.assertEqual(nut['result'],'Yes')
        self.assertEqual(bolt['comments'],['Bolt note'])
        self.assertEqual(nut['comments'],['Nut CCC note'])
        self.assertFalse(bolt['substrateFinishUsed'])
        self.assertTrue(nut['substrateFinishUsed'])
        self.item['finish']='TSA'
        results=self.db.compatibility(self.query)[1]['results']
        self.assertEqual(results[0]['result'],'No')
        self.assertEqual(results[1]['result'],'No')
        self.assertEqual(extra['status'],'missing_values')

    def test_additional_component_uses_cg_ch_and_ignores_substrate_finish(self):
        self.row[84:86]=['Titanium alloy','None']
        self.item['finish']='Unlisted finish'
        results=self.db.compatibility(self.query)[1]['results']
        self.assertEqual(results[2]['status'],'matched')
        self.assertEqual(results[2]['result'],'No')
        self.assertEqual(results[2]['comments'],['Bolt note'])
        self.assertEqual(results[1]['status'],'no_match')

    def test_checkbox_and_five_substrates(self):
        self.item['nutSide']=False
        self.query['substrates']=[dict(self.item) for _ in range(5)]
        results=self.db.compatibility(self.query)[1]['results']
        self.assertEqual(len(results),10)
        self.assertEqual({r['substrateIndex'] for r in results},{1,2,3,4,5})
        self.assertFalse(any('Nut /' in r['component'] for r in results))
        self.query['substrates'].append(dict(self.item))
        with self.assertRaises(ValueError):self.db.compatibility(self.query)

    def test_missing_no_match_conflict(self):
        self.row[80]=None
        self.assertEqual(self.db.compatibility(self.query)[1]['results'][0]['status'],'missing_values')
        self.row[80]='Unknown'
        self.assertEqual(self.db.compatibility(self.query)[1]['results'][0]['status'],'no_match')
        self.row[80]='Titanium alloy'
        self.db.galvanic.append(['Titanium alloy','None','Aluminum alloy','Yes','Other note'])
        result=self.db.compatibility(self.query)[1]['results'][0]
        self.assertEqual(result['status'],'conflict')
        self.assertEqual(result['comments'],['Bolt note','Other note'])

    def test_yes_invalid_inputs_and_unknown_part(self):
        self.query['mode']='yes'
        self.assertEqual(self.db.compatibility(self.query)[1]['status'],'complete')
        self.query['part']='unknown'
        self.assertEqual(self.db.compatibility(self.query)[0],404)
        for payload in [None,{},dict(self.query,substrates=[]),dict(self.query,mode='invalid'),
                        dict(self.query,substrates=[dict(self.item,finish='')]),
                        dict(self.query,substrates=[dict(self.item,nutSide='false')])]:
            with self.assertRaises(ValueError):self.db.compatibility(payload)

    def test_yes_uses_bonding_tables_and_no_keeps_original_tables(self):
        self.row[84:86]=['Titanium alloy','None']
        no=self.db.compatibility(self.query)[1]['results']
        self.query['mode']='yes'
        yes=self.db.compatibility(self.query)[1]['results']
        self.assertEqual([r['result'] for r in no],['No','Yes','No'])
        self.assertEqual([r['result'] for r in yes],['Yes','No','Yes'])
        self.assertEqual(yes[0]['comments'],['Bonding bolt note'])
        self.assertEqual(yes[2]['comments'],['Bonding bolt note'])
        self.assertEqual(yes[1]['comments'],[])
        self.item['finish']='Other finish'
        results=self.db.compatibility(self.query)[1]['results']
        self.assertEqual(results[0]['result'],'Yes')
        self.assertEqual(results[1]['status'],'no_match')
        self.item['nutSide']=False
        self.query['substrates']=[dict(self.item) for _ in range(5)]
        results=self.db.compatibility(self.query)[1]['results']
        self.assertEqual(len(results),10)
        self.assertTrue(all(r['result']=='Yes' for r in results))

    def test_yes_has_no_fallback_to_no_table(self):
        self.query['mode']='yes'
        self.db.bonding_galvanic=[]
        self.db.bonding_nut_galvanic=[]
        results=self.db.compatibility(self.query)[1]['results']
        self.assertEqual(results[0]['status'],'no_match')
        self.assertEqual(results[1]['status'],'no_match')

    def test_bonding_column_extraction(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'test.xlsx'
            book=openpyxl.Workbook();sheet=book.active;sheet.title='Database'
            sheet.append([None]*86);sheet.append(['Header']*86);sheet.append(self.row)
            table=book.create_sheet('Galvanic compatability');table.append(['Header']*32)
            row=[None]*32
            row[:5]=['old','old','old','No','Old note']
            row[20:25]=['Titanium alloy','None','Aluminum alloy','Yes','Y comment']
            row[26:32]=['Titanium alloy','None','Aluminum alloy','CCC','No','AF must not be used']
            table.append(row);book.save(path);book.close()
            db=server.Database(path)
            self.query['mode']='yes'
            results=db.compatibility(self.query)[1]['results']
            self.assertEqual(results[0]['result'],'Yes')
            self.assertEqual(results[0]['comments'],['Y comment'])
            self.assertEqual(results[1]['result'],'No')
            self.assertEqual(results[1]['comments'],[])

    def test_pair_result_both_modes_independent_of_checkbox(self):
        self.item['nutSide']=False
        self.query['substrates']=[dict(self.item) for _ in range(5)]
        result=self.db.compatibility(self.query)[1]
        self.assertEqual(len(result['results']),10)
        self.assertEqual(result['pairResult']['result'],'Yes')
        self.assertEqual(result['pairResult']['comments'],['Pair note'])
        self.assertEqual(result['pairResult']['values'],['Titanium alloy','None','Titanium alloy','None'])
        self.query['mode']='yes'
        pair=self.db.compatibility(self.query)[1]['pairResult']
        self.assertEqual(pair['result'],'No')
        self.assertEqual(pair['comments'],['Bonding pair note'])

    def test_pair_column_order_and_failure_cases(self):
        self.row[80:84]=['Bolt material','Bolt finish','Nut material','Nut finish']
        self.db.pair_galvanic=[['Bolt material','Bolt finish','Nut material','Nut finish','No','Pair-specific note']]
        pair=self.db.compatibility(self.query)[1]['pairResult']
        self.assertEqual(pair['result'],'No')
        self.assertEqual(pair['comments'],['Pair-specific note'])
        self.db.pair_galvanic.append(['Bolt material','Bolt finish','Nut material','Nut finish','Yes','Conflict note'])
        self.assertEqual(self.db.compatibility(self.query)[1]['pairResult']['status'],'conflict')
        self.row[81]='Unlisted finish'
        self.assertEqual(self.db.compatibility(self.query)[1]['pairResult']['status'],'no_match')
        self.row[81]=None
        self.assertEqual(self.db.compatibility(self.query)[1]['pairResult']['status'],'missing_values')

    def test_yes_pair_no_match_missing_and_conflict(self):
        self.query['mode']='yes'
        self.db.bonding_pair_galvanic=[]
        self.assertEqual(self.db.compatibility(self.query)[1]['pairResult']['status'],'no_match')
        keys=['Titanium alloy','None','Titanium alloy','None']
        self.db.bonding_pair_galvanic=[keys+['Yes','First'],keys+['No','Second']]
        self.assertEqual(self.db.compatibility(self.query)[1]['pairResult']['status'],'conflict')
        self.row[82]=None
        self.assertEqual(self.db.compatibility(self.query)[1]['pairResult']['status'],'missing_values')

    def test_yes_pair_excel_column_positions(self):
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'test.xlsx'
            book=openpyxl.Workbook();sheet=book.active;sheet.title='Database'
            sheet.append([None]*86);sheet.append(['Header']*86);sheet.append(self.row)
            table=book.create_sheet('Galvanic compatability');table.append(['Header']*39)
            row=[None]*39
            row[13:19]=['Titanium alloy','None','Titanium alloy','None','Yes','S note']
            row[33:39]=['Titanium alloy','None','Titanium alloy','None','No','AM note']
            table.append(row);book.save(path);book.close()
            db=server.Database(path)
            self.query['mode']='yes'
            pair=db.compatibility(self.query)[1]['pairResult']
            self.assertEqual(pair['result'],'No');self.assertEqual(pair['comments'],['AM note'])
            self.query['mode']='no'
            pair=db.compatibility(self.query)[1]['pairResult']
            self.assertEqual(pair['result'],'Yes');self.assertEqual(pair['comments'],['S note'])

    def test_real_workbook_options(self):
        db=server.Database(ROOT/'data/blindfastener.xlsx')
        meta=db.metadata()
        self.assertIn('CCC',meta['substrateFinishes'])
        self.assertIn('None',meta['substrateFinishes'])
        self.assertIn('Aluminum alloy',meta['substrates'])
        self.assertEqual(len(meta['partNumbers']),2640)

    def test_post_endpoint(self):
        http=server.ThreadingHTTPServer(('127.0.0.1',0),server.make_handler(self.db))
        worker=threading.Thread(target=http.serve_forever,daemon=True);worker.start()
        url=f'http://127.0.0.1:{http.server_port}/api/compatibility'
        try:
            request=Request(url,data=json.dumps(self.query).encode(),headers={'Content-Type':'application/json'})
            with urlopen(request) as response:
                self.assertEqual(json.load(response)['results'][1]['result'],'Yes')
            with self.assertRaises(HTTPError) as error:
                urlopen(Request(url,data=b'{invalid'))
            self.assertEqual(error.exception.code,400)
        finally:
            http.shutdown();http.server_close();worker.join()

if __name__=='__main__':unittest.main()
